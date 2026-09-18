import unicodedata
from pathlib import Path

import xlrd
from openpyxl import load_workbook


def normalizar_texto(valor) -> str:
    """Normaliza texto para facilitar comparaciones."""

    texto = str(valor).strip().lower()

    texto = unicodedata.normalize("NFD", texto)

    texto = "".join(
        caracter
        for caracter in texto
        if unicodedata.category(caracter) != "Mn"
    )

    return texto


def normalizar_codigo(codigo) -> str:
    """Normaliza un código de artículo."""

    codigo = str(codigo).strip()

    if codigo.upper().startswith("ALT-"):
        codigo = codigo[4:]

    return codigo


def normalizar_cantidad(cantidad) -> int:
    """Convierte una cantidad del Excel a entero."""

    if isinstance(cantidad, bool):
        raise ValueError(
            f"Cantidad inválida: {cantidad!r}"
        )

    if isinstance(cantidad, (int, float)):
        if float(cantidad).is_integer():
            return int(cantidad)

    texto = str(cantidad).strip()

    try:
        numero = float(
            texto.replace(",", ".")
        )

        if numero.is_integer():
            return int(numero)

    except ValueError:
        pass

    raise ValueError(
        f"Cantidad inválida: {cantidad!r}"
    )


def encontrar_encabezado(filas):
    """
    Encuentra las columnas Código y Cantidad.

    Devuelve:
        (fila_encabezado, columna_codigo, columna_cantidad)
    """

    for numero_fila, fila in enumerate(
        filas,
        start=1
    ):
        columnas = {}

        for numero_columna, valor in enumerate(
            fila,
            start=1
        ):
            if valor is None:
                continue

            encabezado = normalizar_texto(
                valor
            )

            if encabezado == "codigo":
                columnas["codigo"] = numero_columna

            elif encabezado == "cantidad":
                columnas["cantidad"] = numero_columna

        if (
            "codigo" in columnas
            and "cantidad" in columnas
        ):
            return (
                numero_fila,
                columnas["codigo"],
                columnas["cantidad"],
            )

    raise ValueError(
        "No se encontró un encabezado "
        "con Código y Cantidad."
    )


def extraer_pedido(
    filas,
    fila_encabezado,
    columna_codigo,
    columna_cantidad,
):
    """
    Extrae los artículos a partir de las columnas detectadas.

    Las filas incompletas o con datos inválidos se ignoran.
    """

    pedido = []

    indice_codigo = columna_codigo - 1
    indice_cantidad = columna_cantidad - 1

    for fila in filas[fila_encabezado:]:
        # Algunas filas de Excel pueden tener menos columnas
        # que el encabezado. Las ignoramos en lugar de provocar
        # un IndexError.
        if len(fila) <= max(
            indice_codigo,
            indice_cantidad,
        ):
            continue

        codigo = fila[indice_codigo]
        cantidad = fila[indice_cantidad]

        if codigo is None or cantidad is None:
            continue

        if str(codigo).strip() == "":
            continue

        codigo = normalizar_codigo(
            codigo
        )

        if not codigo:
            continue

        try:
            cantidad = normalizar_cantidad(
                cantidad
            )
        except ValueError:
            continue

        if cantidad <= 0:
            continue

        pedido.append(
            {
                "codigo": codigo,
                "cantidad": cantidad,
            }
        )

    return pedido


def leer_xlsx(ruta_archivo: str) -> list[dict]:
    """Lee un archivo .xlsx."""

    workbook = load_workbook(
        ruta_archivo,
        read_only=True,
        data_only=True,
    )

    try:
        ws = workbook[
            workbook.sheetnames[0]
        ]

        filas = list(
            ws.iter_rows(
                values_only=True
            )
        )

        (
            fila_encabezado,
            columna_codigo,
            columna_cantidad,
        ) = encontrar_encabezado(
            filas
        )

        return extraer_pedido(
            filas,
            fila_encabezado,
            columna_codigo,
            columna_cantidad,
        )

    finally:
        workbook.close()


def leer_xls(ruta_archivo: str) -> list[dict]:
    """Lee un archivo .xls."""

    workbook = xlrd.open_workbook(
        ruta_archivo
    )

    sheet = workbook.sheet_by_index(0)

    filas = []

    for numero_fila in range(
        sheet.nrows
    ):
        filas.append(
            sheet.row_values(
                numero_fila
            )
        )

    (
        fila_encabezado,
        columna_codigo,
        columna_cantidad,
    ) = encontrar_encabezado(
        filas
    )

    return extraer_pedido(
        filas,
        fila_encabezado,
        columna_codigo,
        columna_cantidad,
    )


def leer_pedido(
    ruta_archivo: str
) -> list[dict]:
    """Lee un pedido .xls o .xlsx."""

    extension = Path(
        ruta_archivo
    ).suffix.lower()

    if extension == ".xls":
        return leer_xls(
            ruta_archivo
        )

    if extension == ".xlsx":
        return leer_xlsx(
            ruta_archivo
        )

    raise ValueError(
        f"Formato no soportado: {extension}. "
        "Utilizá .xls o .xlsx."
    )


if __name__ == "__main__":
    archivo = input(
        "Ruta del archivo Excel: "
    ).strip()

    pedido = leer_pedido(
        archivo
    )

    print()
    print(
        f"Artículos encontrados: "
        f"{len(pedido)}"
    )
    print()

    for articulo in pedido[:10]:
        print(
            articulo["codigo"],
            "x",
            articulo["cantidad"],
        )